"""
Bulk listing import/export via a two-sheet XLSX workbook (Listings +
HotelRooms, linked by a per-file `row_id` — not a database ID). Import reuses
ListingSerializer/HotelRoomSerializer directly so bulk-created listings obey
the exact same validation (required fields, minimum price, etc.) as one
created through the normal wizard.

Scope, deliberately: this creates NEW listings only (no bulk-edit of existing
ones — re-importing never updates a row by ID), and never touches images —
downloading arbitrary image URLs server-side is an SSRF risk (see
docs/SECURITY_TESTING.md TEST-INPUT-06), and XLSX cells can't carry binary
image data in a way worth building a mini image-hosting pipeline around.
Bulk-imported listings land at `pending_review` exactly like a normal
submission — the host/admin still adds photos and completes ownership
verification per listing afterward (the "resumable draft" flow already
handles a listing sitting unverified — see CreateListing.tsx).
"""
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from decimal import Decimal, InvalidOperation

from .models import Listing, HotelRoom
from .serializers import ListingSerializer, HotelRoomSerializer

LISTING_SHEET = 'Listings'
ROOMS_SHEET = 'HotelRooms'
INSTRUCTIONS_SHEET = 'Instructions'

# (header, field name, kind) — kind drives both writing and parsing.
# 'list' fields are comma-separated in the sheet.
LISTING_COLUMNS = [
    ('row_id', 'row_id', 'int'),
    ('owner_email', 'owner_email', 'str'),
    ('title', 'title', 'str'),
    ('description', 'description', 'str'),
    ('price', 'price', 'decimal'),
    ('property_type', 'property_type', 'str'),
    ('privacy_type', 'privacy_type', 'str'),
    ('address', 'address', 'str'),
    ('city', 'city', 'str'),
    ('state', 'state', 'str'),
    ('country', 'country', 'str'),
    ('latitude', 'latitude', 'decimal'),
    ('longitude', 'longitude', 'decimal'),
    ('check_in_time', 'check_in_time', 'str'),
    ('check_out_time', 'check_out_time', 'str'),
    ('self_checkin', 'self_checkin', 'bool'),
    ('square_footage', 'square_footage', 'int'),
    ('bedrooms', 'bedrooms', 'int'),
    ('beds', 'beds', 'int'),
    ('bathrooms', 'bathrooms', 'int'),
    ('max_guests', 'max_guests', 'int'),
    ('amenities', 'amenities', 'list'),
    ('highlights', 'highlights', 'list'),
    ('booking_mode', 'booking_mode', 'str'),
    ('cancellation_policy', 'cancellation_policy', 'str'),
    ('weekend_premium_percent', 'weekend_premium_percent', 'int'),
    ('new_listing_promo', 'new_listing_promo', 'bool'),
    ('last_minute_discount_enabled', 'last_minute_discount_enabled', 'bool'),
    ('last_minute_discount_percent', 'last_minute_discount_percent', 'int'),
    ('weekly_discount_enabled', 'weekly_discount_enabled', 'bool'),
    ('weekly_discount_percent', 'weekly_discount_percent', 'int'),
    ('monthly_discount_enabled', 'monthly_discount_enabled', 'bool'),
    ('monthly_discount_percent', 'monthly_discount_percent', 'int'),
    ('exterior_camera', 'exterior_camera', 'bool'),
    ('noise_monitor', 'noise_monitor', 'bool'),
    ('weapons_on_property', 'weapons_on_property', 'bool'),
    ('pricing_type', 'pricing_type', 'str'),
    ('payment_schedule', 'payment_schedule', 'str'),
    ('lease_term_months', 'lease_term_months', 'int'),
]

ROOM_COLUMNS = [
    ('row_id', 'row_id', 'int'),
    ('name', 'name', 'str'),
    ('room_type', 'room_type', 'str'),
    ('description', 'description', 'str'),
    ('price_per_night', 'price_per_night', 'decimal'),
    ('max_occupancy', 'max_occupancy', 'int'),
    ('beds', 'beds', 'int'),
    ('bed_type', 'bed_type', 'str'),
    ('bathrooms', 'bathrooms', 'int'),
    ('amenities', 'amenities', 'list'),
    ('total_count', 'total_count', 'int'),
]

EXAMPLE_LISTING_ROW = {
    'row_id': 1, 'owner_email': '', 'title': 'Cozy Two Bedroom Apartment',
    'description': 'A bright, quiet apartment close to downtown.',
    'price': 45, 'property_type': 'apartment', 'privacy_type': 'entire_place',
    'address': '12 Randall St', 'city': 'Monrovia', 'state': 'Montserrado', 'country': 'Liberia',
    'latitude': '', 'longitude': '', 'check_in_time': '15:00', 'check_out_time': '11:00',
    'self_checkin': False, 'square_footage': 900, 'bedrooms': 2, 'beds': 2, 'bathrooms': 1,
    'max_guests': 4, 'amenities': 'wifi, kitchen, parking', 'highlights': 'quiet street',
    'booking_mode': 'approve_first', 'cancellation_policy': 'flexible',
    'weekend_premium_percent': 0, 'new_listing_promo': False,
    'last_minute_discount_enabled': False, 'last_minute_discount_percent': 0,
    'weekly_discount_enabled': False, 'weekly_discount_percent': 0,
    'monthly_discount_enabled': False, 'monthly_discount_percent': 0,
    'exterior_camera': False, 'noise_monitor': False, 'weapons_on_property': False,
    'pricing_type': 'nightly', 'payment_schedule': '', 'lease_term_months': '',
}

EXAMPLE_ROOM_ROW = {
    'row_id': 1, 'name': 'Standard Queen', 'room_type': 'standard', 'description': '',
    'price_per_night': 30, 'max_occupancy': 2, 'beds': 1, 'bed_type': 'queen',
    'bathrooms': 1, 'amenities': 'wifi, ac', 'total_count': 3,
}

INSTRUCTIONS_TEXT = [
    "How to use this workbook",
    "",
    "1. Fill in one row per listing on the 'Listings' sheet. row_id is a "
    "number you choose to identify the row within THIS file only (not a "
    "database ID) — it's how a listing links to its room types below.",
    "2. Only fill in the 'HotelRooms' sheet for listings whose property_type "
    "is 'hotels' or 'lodge' — every row there must reference a row_id that "
    "exists on the Listings sheet. Leave it empty for other property types.",
    "3. amenities and highlights are comma-separated (e.g. 'wifi, kitchen, parking').",
    "4. True/false columns accept TRUE/FALSE, yes/no, or 1/0.",
    "5. owner_email: leave blank to import as your own listings. Admins "
    "importing on behalf of another host/agent must fill this in with that "
    "user's account email.",
    "6. Imported listings are created exactly like one made through the "
    "listing wizard (same required fields, same minimum price) and land as "
    "'Pending Review' — NOT published yet. Photos aren't part of this "
    "import; add them per listing afterward. Each listing still needs its "
    "ownership verification completed before it can go live, same as any "
    "other listing.",
    "7. Re-uploading this file does not update existing listings — every "
    "import creates brand-new listings.",
]


def _cell(kind, value):
    if value is None:
        return ''
    if kind == 'bool':
        return 'TRUE' if value else 'FALSE'
    if kind == 'list':
        return ', '.join(value) if isinstance(value, (list, tuple)) else (value or '')
    return value


def _parse_cell(kind, value):
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if kind == 'int':
        return int(Decimal(str(value)))
    if kind == 'decimal':
        return Decimal(str(value))
    if kind == 'bool':
        return str(value).strip().lower() in ('true', 'yes', '1', 'y')
    if kind == 'list':
        return [v.strip() for v in str(value).split(',') if v.strip()]
    return str(value).strip()


def _write_sheet(ws: Worksheet, columns, rows):
    ws.append([header for header, _, _ in columns])
    for row in rows:
        ws.append([_cell(kind, row.get(field)) for _, field, kind in columns])


def build_template_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb.create_sheet(LISTING_SHEET), LISTING_COLUMNS, [EXAMPLE_LISTING_ROW])
    _write_sheet(wb.create_sheet(ROOMS_SHEET), ROOM_COLUMNS, [EXAMPLE_ROOM_ROW])
    instructions = wb.create_sheet(INSTRUCTIONS_SHEET, 0)
    for line in INSTRUCTIONS_TEXT:
        instructions.append([line])
    instructions.column_dimensions['A'].width = 100
    return wb


def build_export_workbook(listings_qs) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    listing_rows = []
    room_rows = []
    for i, listing in enumerate(listings_qs.prefetch_related('hotel_rooms'), start=1):
        row = {'row_id': i, 'owner_email': listing.owner.email}
        for _, field, _kind in LISTING_COLUMNS:
            if field in ('row_id', 'owner_email'):
                continue
            row[field] = getattr(listing, field)
        listing_rows.append(row)
        for room in listing.hotel_rooms.all():
            room_row = {'row_id': i}
            for _, field, _kind in ROOM_COLUMNS:
                if field == 'row_id':
                    continue
                room_row[field] = getattr(room, field)
            room_rows.append(room_row)
    _write_sheet(wb.create_sheet(LISTING_SHEET), LISTING_COLUMNS, listing_rows)
    _write_sheet(wb.create_sheet(ROOMS_SHEET), ROOM_COLUMNS, room_rows)
    return wb


def _read_sheet_rows(ws: Worksheet, columns):
    if ws is None:
        return []
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        return []
    header_index = {h: i for i, h in enumerate(header_row) if h}
    rows = []
    for excel_row_number, raw_row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == '' for v in raw_row):
            continue  # skip blank rows
        parsed = {}
        for header, field, kind in columns:
            idx = header_index.get(header)
            raw_value = raw_row[idx] if idx is not None and idx < len(raw_row) else None
            try:
                value = _parse_cell(kind, raw_value)
            except (InvalidOperation, ValueError):
                value = None
            # A blank cell means "use the field's own default" — including the
            # key with an explicit None would instead fail validation for
            # optional-but-non-nullable fields (e.g. check_in_time, which has
            # a default but no null=True). row_id is the one exception: it's
            # popped separately by the caller and never reaches the serializer.
            if value is not None or field == 'row_id':
                parsed[field] = value
        rows.append((excel_row_number, parsed))
    return rows


def import_workbook(file_obj, requesting_user, is_admin: bool):
    """
    Returns (created_listings, row_errors) where row_errors is a list of
    {'row': <excel row number>, 'errors': {...}} for rows that failed
    validation — valid rows still import even if others fail.
    """
    wb = openpyxl.load_workbook(file_obj, data_only=True)
    listing_rows = _read_sheet_rows(wb[LISTING_SHEET] if LISTING_SHEET in wb.sheetnames else None, LISTING_COLUMNS)
    room_rows = _read_sheet_rows(wb[ROOMS_SHEET] if ROOMS_SHEET in wb.sheetnames else None, ROOM_COLUMNS)

    rooms_by_row_id = {}
    for _excel_row, room_data in room_rows:
        rooms_by_row_id.setdefault(room_data.get('row_id'), []).append(room_data)

    User = requesting_user.__class__
    created_listings = []
    row_errors = []

    for excel_row_number, data in listing_rows:
        row_id = data.pop('row_id', None)
        owner_email = (data.pop('owner_email', '') or '').strip()

        if owner_email:
            if not is_admin:
                row_errors.append({'row': excel_row_number, 'errors': {
                    'owner_email': 'Only admins can import listings on behalf of another account.'
                }})
                continue
            owner = User.objects.filter(email__iexact=owner_email).first()
            if not owner:
                row_errors.append({'row': excel_row_number, 'errors': {
                    'owner_email': f'No account found for {owner_email}.'
                }})
                continue
        elif is_admin:
            row_errors.append({'row': excel_row_number, 'errors': {
                'owner_email': 'Required for an admin-driven import — whose listing is this?'
            }})
            continue
        else:
            owner = requesting_user

        serializer = ListingSerializer(data=data)
        if not serializer.is_valid():
            row_errors.append({'row': excel_row_number, 'errors': serializer.errors})
            continue

        listing = serializer.save(owner=owner, status='pending_review')

        for room_data in rooms_by_row_id.get(row_id, []):
            room_data = dict(room_data)
            room_data.pop('row_id', None)
            room_data['listing'] = listing.id
            room_serializer = HotelRoomSerializer(data=room_data)
            if room_serializer.is_valid():
                room_serializer.save(listing=listing)
            else:
                row_errors.append({'row': excel_row_number, 'errors': {
                    'hotel_rooms': f'A room type failed validation and was skipped: {room_serializer.errors}'
                }})

        created_listings.append(listing)

    return created_listings, row_errors
